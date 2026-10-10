import csv
import io
import json
import time
from datetime import datetime, timezone
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, Query, UploadFile
from fastapi.responses import PlainTextResponse, StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.audit import EVENT_NAMES, audit
from app.config import get_settings
from app.db import get_db
from app.domain import PROCEDURE_NAMES, VERDICT_NAMES, CanonicalTender, Preferences
from app.egrul.providers import CompanyNotFound, EgrulError, FnsProvider, InvalidInn, get_company, try_get_company
from app.eis.parser import ParseError, parse_bytes
from app.models import AuditLog, Batch, BatchFile, Feedback, Profile, ProfileVersion, Score, Tender
from app.nlp import llm
from app.nlp.criteria import parse_criteria
from app.reference.regions import DISTRICTS, REGIONS
from app.scoring import ahp
from app.scoring.factors import FACTORS
from app import services
from app.worker import dispatch_batch, unpack

router = APIRouter(prefix="/api")
settings = get_settings()
EVAL_REPORT = Path(__file__).resolve().parents[1] / "eval" / "report.json"


# ---------- служебное ----------

@router.get("/health", summary="Состояние сервиса и зависимостей")
def health(db: Session = Depends(get_db)):
    db.execute(select(1))
    return {
        "status": "ok",
        "llm": llm.provider_status(),
        "egrul_provider": "dadata" if settings.egrul_provider == "dadata" and settings.dadata_api_key else "mock",
        "egrul_chain": (["dadata"] if settings.egrul_provider == "dadata" and settings.dadata_api_key else [])
                       + ["mock"] + (["fns"] if settings.egrul_public_fns else []),
        "fns_breaker": FnsProvider.breaker.state(),
        "queue": "celery" if settings.redis_url else "in-process",
        "eis_api": bool(settings.eis_token),
        "llm_review": settings.llm_review,
    }


@router.get("/factors", summary="Схема факторов для панели настроек и паутинки")
def factors():
    return {
        "factors": [
            {"key": f.key, "label": f.label, "base_weight": f.base_weight, "description": f.description, "controls": f.controls}
            for f in FACTORS
        ],
        "regions": [{"code": r.code, "name": r.name, "district": r.district} for r in REGIONS],
        "districts": DISTRICTS,
        "procedures": PROCEDURE_NAMES,
        "verdicts": VERDICT_NAMES,
        "defaults": json.loads(Preferences().model_dump_json()),
    }


# ---------- компании (ЕГРЮЛ) ----------

@router.get("/companies/{inn}", summary="Карточка организации из ЕГРЮЛ по ИНН")
def company(inn: str, refresh: bool = False, db: Session = Depends(get_db)):
    try:
        return get_company(db, inn, refresh=refresh)
    except InvalidInn as exc:
        raise HTTPException(400, str(exc))
    except CompanyNotFound as exc:
        raise HTTPException(404, str(exc))
    except EgrulError as exc:
        raise HTTPException(502, str(exc))


# ---------- профили ----------

@router.get("/fuzzy/curves", summary="Функции принадлежности нечётких переменных для графиков (по профилю)")
def fuzzy_curves(profile_id: int | None = None, db: Session = Depends(get_db)):
    from app.scoring import fuzzy

    prefs = Preferences.model_validate(services.current_version(db, _profile(db, profile_id)).preferences)
    return fuzzy.curves(prefs.price.min_rub, prefs.price.max_rub, prefs.timing.min_days_to_deadline)


class PairIn(BaseModel):
    a: str
    b: str
    value: float = Field(gt=0, description="«a важнее b в value раз» по шкале Саати 1/9…9")


class WeightsRequest(BaseModel):
    method: str = Field("roc", pattern="^(roc|ahp)$")
    ranking: list[str] = Field(default_factory=list, description="для roc: факторы от важного к неважному")
    pairs: list[PairIn] = Field(default_factory=list, description="для ahp: попарные сравнения")


@router.get("/weights/questions", summary="Вопросы мастера весов: пары факторов для МАИ")
def weight_questions():
    labels = {f.key: f.label for f in FACTORS}
    return {"pairs": [{"a": a, "b": b, "a_label": labels[a], "b_label": labels[b]} for a, b in ahp.DEFAULT_PAIRS],
            "factors": [{"key": f.key, "label": f.label, "base_weight": f.base_weight} for f in FACTORS]}


@router.post("/weights", summary="Веса факторов: ранжирование (ROC) или попарные сравнения (МАИ Саати)")
def weights(req: WeightsRequest):
    keys = [f.key for f in FACTORS]
    labels = {f.key: f.label for f in FACTORS}
    if req.method == "roc":
        ranking = [k for k in req.ranking if k in labels]
        if sorted(ranking) != sorted(keys):
            raise HTTPException(422, "В ранжировании должны быть все факторы ровно по одному разу")
        return {"method": "roc", "weights": ahp.roc_weights(ranking), "cr": None, "consistent": True, "advice": []}
    if len(req.pairs) < len(keys) - 1:
        raise HTTPException(422, f"Нужно хотя бы {len(keys) - 1} сравнений, чтобы связать все факторы")
    pairs = [(p.a, p.b, p.value) for p in req.pairs if p.a in labels and p.b in labels and p.a != p.b]
    idx = {k: i for i, k in enumerate(keys)}
    loose = [labels[k] for k in keys if not ahp.connected(idx, pairs, keys[0], k)]
    if loose:
        raise HTTPException(422, "Факторы не связаны сравнениями с остальными: " + ", ".join(loose))
    res = ahp.weights_from_pairs(keys, pairs)
    return {"method": "ahp", "weights": res["weights"], "cr": res["cr"], "consistent": res["consistent"],
            "lambda_max": res["lambda_max"], "worst_pairs": res["worst_pairs"], "advice": ahp.advise(res, labels)}


class ProfileCreate(BaseModel):
    name: str = Field(..., min_length=1, max_length=200)
    company_inn: str
    criteria_text: str | None = None
    preferences: Preferences | None = None


class ProfileUpdate(BaseModel):
    preferences: Preferences
    criteria_text: str | None = None


class ParseRequest(BaseModel):
    text: str = Field(..., max_length=5000)
    use_llm: bool = True


def _profile(db: Session, profile_id: int | None) -> Profile:
    profile = db.get(Profile, profile_id) if profile_id else db.scalar(select(Profile).order_by(Profile.id))
    if profile is None:
        raise HTTPException(404, "Профиль не найден — создайте его через POST /api/profiles")
    return profile


def _profile_out(db: Session, profile: Profile) -> dict:
    pv = services.current_version(db, profile)
    company_card = services.profile_company(db, profile)
    return {
        "id": profile.id,
        "name": profile.name,
        "company_inn": profile.company_inn,
        "company": company_card,
        "version": pv.version,
        "criteria_text": pv.criteria_text,
        "preferences": pv.preferences,
        "versions": [{"version": v.version, "created_at": v.created_at} for v in profile.versions],
    }


@router.get("/profiles", summary="Список профилей")
def list_profiles(db: Session = Depends(get_db)):
    return [{"id": p.id, "name": p.name, "company_inn": p.company_inn, "version": p.current_version}
            for p in db.scalars(select(Profile).order_by(Profile.id))]


@router.post("/profiles", summary="Создать профиль (критерии текстом и/или готовые настройки)")
def create_profile(body: ProfileCreate, db: Session = Depends(get_db)):
    prefs = body.preferences or Preferences()
    parsed = None
    if body.criteria_text:
        parsed = parse_criteria(body.criteria_text, prefs)
        prefs = parsed.preferences
    try:
        profile = services.create_profile(db, body.name, body.company_inn, prefs, body.criteria_text)
    except InvalidInn as exc:
        raise HTTPException(400, str(exc))
    except EgrulError as exc:
        raise HTTPException(404, str(exc))
    services.rescore_all(db, profile)
    out = _profile_out(db, profile)
    out["parsed"] = parsed.to_dict() if parsed else None
    return out


@router.get("/profiles/{profile_id}", summary="Профиль с текущими настройками")
def get_profile(profile_id: int, db: Session = Depends(get_db)):
    return _profile_out(db, _profile(db, profile_id))


@router.post("/profiles/{profile_id}/parse", summary="Разобрать текст критериев в настройки (без сохранения)")
def parse_profile_text(profile_id: int, body: ParseRequest, db: Session = Depends(get_db)):
    profile = _profile(db, profile_id)
    base = Preferences()
    base.geo.same_district_score = Preferences.model_validate(services.current_version(db, profile).preferences).geo.same_district_score
    outcome = parse_criteria(body.text, base, use_llm=body.use_llm, timeout=settings.llm_parse_timeout_seconds)
    audit("criteria_parsed", f"Распознано правил: {len(outcome.recognized)}, не понято фраз: {len(outcome.unparsed)}",
          duration_ms=outcome.elapsed_ms, engine=outcome.engine, unparsed=outcome.unparsed[:5])
    return outcome.to_dict()


@router.put("/profiles/{profile_id}", summary="Сохранить настройки как новую версию и пересчитать ленту")
def update_profile(profile_id: int, body: ProfileUpdate, db: Session = Depends(get_db)):
    profile = _profile(db, profile_id)
    started = time.perf_counter()
    services.new_version(db, profile, body.preferences, body.criteria_text)
    n = services.rescore_all(db, profile)
    audit("profile_saved", f"Профиль «{profile.name}» → версия {profile.current_version}, пересчитано {n} закупок",
          duration_ms=(time.perf_counter() - started) * 1000, profile_id=profile.id, version=profile.current_version)
    out = _profile_out(db, profile)
    out["rescored"] = n
    return out


# ---------- закупки ----------

async def _read_upload(file: UploadFile) -> bytes:
    content = await file.read()
    if len(content) > settings.max_upload_mb * 1024 * 1024:
        raise HTTPException(413, f"Файл больше {settings.max_upload_mb} МБ")
    if not content.strip():
        raise HTTPException(400, "Пустой файл")
    return content


@router.post("/tenders/parse", summary="Разобрать выгрузку ЕИС без сохранения")
async def parse_tender(file: UploadFile = File(...)) -> CanonicalTender:
    content = await _read_upload(file)
    try:
        return parse_bytes(content, file.filename or "")
    except ParseError as exc:
        raise HTTPException(422, str(exc))


@router.post("/tenders", summary="Загрузить извещение (XML/JSON) и сразу оценить")
async def upload_tender(file: UploadFile = File(...), profile_id: int | None = Form(None), db: Session = Depends(get_db)):
    content = await _read_upload(file)
    profile = _profile(db, profile_id)
    started = time.perf_counter()
    try:
        row = services.ingest(db, file.filename or "upload.xml", content)
    except ParseError as exc:
        audit("error", f"Файл {file.filename} не разобран — {exc}", level="warning")
        raise HTTPException(422, str(exc))
    pv = services.current_version(db, profile)
    score = services.score_row(db, row, pv, services.profile_company(db, profile), force=True)
    audit("tender_uploaded", f"{row.purchase_number}: {score.score:.0f}% — {score.result.get('verdict_name')}",
          duration_ms=(time.perf_counter() - started) * 1000, purchase_number=row.purchase_number, filename=file.filename)
    return {"tender_id": row.id, "score_id": score.id, "result": score.result}


@router.get("/eis/notice/{reestr_number}", response_class=PlainTextResponse,
            summary="XML извещения 44-ФЗ из ЕИС по номеру (для формы быстрой оценки)")
def eis_notice(reestr_number: str):
    from app.eis import public
    from app.eis.getdocs import EisApiError

    if not reestr_number.isdigit() or len(reestr_number) != 19:
        raise HTTPException(422, "Номер закупки 44-ФЗ — 19 цифр")
    try:
        name, content = public.fetch_notice(reestr_number)
    except EisApiError as exc:
        raise HTTPException(502, str(exc))
    return PlainTextResponse(content.decode("utf-8", errors="replace"), media_type="application/xml",
                             headers={"Content-Disposition": f'attachment; filename="{name}"'})


class ByNumberRequest(BaseModel):
    reestr_number: str = Field(pattern=r"^\d{19}$", description="Реестровый номер закупки 44-ФЗ (19 цифр)")
    profile_id: int | None = None


@router.post("/tenders/by-number", summary="Загрузить извещение 44-ФЗ из ЕИС по номеру (печатная форма ЕИС; при токене — API getDocsIP)")
def tender_by_number(req: ByNumberRequest, db: Session = Depends(get_db)):
    from app.eis import public
    from app.eis.getdocs import EisApiError, fetch_notice

    started = time.perf_counter()
    try:
        try:
            filename, content = public.fetch_notice(req.reestr_number)  # открытая печатная форма — без токена
        except EisApiError:
            if not settings.eis_token:
                raise
            filename, content = fetch_notice(req.reestr_number)  # официальный API по токену
        row = services.ingest(db, filename, content)
    except EisApiError as exc:
        audit("error", f"ЕИС по номеру {req.reestr_number}: {exc}", level="warning")
        raise HTTPException(502, str(exc))
    except ParseError as exc:
        raise HTTPException(422, f"Документ из ЕИС не разобран: {exc}")
    profile = _profile(db, req.profile_id)
    score = services.score_row(db, row, services.current_version(db, profile), services.profile_company(db, profile), force=True)
    audit("tender_uploaded", f"{row.purchase_number} из ЕИС: {score.score:.0f}% — {score.result.get('verdict_name')}",
          duration_ms=(time.perf_counter() - started) * 1000, purchase_number=row.purchase_number, source="eis_api")
    return {"tender_id": row.id, "score_id": score.id, "result": score.result}


@router.get("/tenders/{tender_id}", summary="Закупка, её оценка по профилю и похожие закупки")
def get_tender(tender_id: int, profile_id: int | None = None, db: Session = Depends(get_db)):
    row = db.get(Tender, tender_id)
    if row is None:
        raise HTTPException(404, "Закупка не найдена")
    profile = _profile(db, profile_id)
    pv = services.current_version(db, profile)
    score = services.score_row(db, row, pv, services.profile_company(db, profile))
    customer = None
    if row.data.get("customer_inn"):
        from app.egrul.providers import try_get_company
        customer = try_get_company(db, row.data["customer_inn"])
    rank = db.scalar(select(func.count()).select_from(Score).where(
        Score.profile_version_id == pv.id, Score.score > score.score))
    return {
        "tender": row.data,
        "tender_id": row.id,
        "score_id": score.id,
        "result": score.result,
        "rank": (rank or 0) + 1,
        "customer": customer,
        "similar": [services._feed_item(t, services.score_row(db, t, pv, services.profile_company(db, profile)))
                    for t in services.similar(db, row)],
        "profile_version": pv.version,
        "bid_status": _bid_status(db, profile.id, row.id),
    }


@router.post("/tenders/{tender_id}/documents", summary="Приложить проект контракта и ТЗ: риски для поставщика и заказчика")
async def tender_documents(tender_id: int, files: list[UploadFile] = File(...), profile_id: int | None = Form(None),
                           db: Session = Depends(get_db)):
    row = db.get(Tender, tender_id)
    if row is None:
        raise HTTPException(404, "Закупка не найдена")
    started = time.perf_counter()
    docs = [(f.filename or "doc", await _read_upload(f)) for f in files]
    profile = _profile(db, profile_id)
    terms = services.attach_documents(db, row, docs, profile.company_inn)
    score = services.score_row(db, row, services.current_version(db, profile), services.profile_company(db, profile))
    audit("documents", f"{row.purchase_number}: разобрано {len(terms.files)} док., находок {len(terms.findings)}",
          duration_ms=(time.perf_counter() - started) * 1000, purchase_number=row.purchase_number,
          files=terms.files, findings=[x.code for x in terms.findings])
    return {"contract": terms, "result": score.result}


def _bid_status(db: Session, profile_id: int, tender_id: int) -> str | None:
    from app.models import Bid

    bid = db.scalar(select(Bid).where(Bid.profile_id == profile_id, Bid.tender_id == tender_id))
    return bid.status if bid else None


class BidIn(BaseModel):
    status: str | None = Field(None, pattern="^(preparing|submitted|won|lost|declined)$")
    profile_id: int | None = None


@router.put("/tenders/{tender_id}/bid", summary="Отметить участие: готовим / подали / выиграли / не выиграли / отказались")
def set_bid(tender_id: int, body: BidIn, db: Session = Depends(get_db)):
    from app.models import Bid

    row = db.get(Tender, tender_id)
    if row is None:
        raise HTTPException(404, "Закупка не найдена")
    profile = _profile(db, body.profile_id)
    bid = db.scalar(select(Bid).where(Bid.profile_id == profile.id, Bid.tender_id == tender_id))
    if body.status is None:
        if bid:
            db.delete(bid)
    elif bid:
        bid.status, bid.updated_at = body.status, datetime.now(timezone.utc)
    else:
        db.add(Bid(profile_id=profile.id, tender_id=tender_id, status=body.status))
    db.commit()
    rescored = services.rescore_all(db, profile)  # свободный лимит обеспечений изменился для всех закупок
    amount, count = services.committed(db, profile.id)
    audit("bid", f"{row.purchase_number}: {services.BID_STATUSES.get(body.status, 'участие снято')}",
          purchase_number=row.purchase_number, status=body.status)
    return {"status": body.status, "status_name": services.BID_STATUSES.get(body.status),
            "committed": amount, "committed_count": count, "rescored": rescored}


@router.get("/bids", summary="Мои заявки и сколько денег занято в обеспечениях")
def list_bids(profile_id: int | None = None, db: Session = Depends(get_db)):
    from app.models import Bid

    profile = _profile(db, profile_id)
    out = []
    for b in db.scalars(select(Bid).where(Bid.profile_id == profile.id).order_by(Bid.updated_at.desc())):
        row = db.get(Tender, b.tender_id)
        if not row:
            continue
        t = services.tender_model(row)
        out.append({"tender_id": row.id, "purchase_number": row.purchase_number, "subject": row.subject,
                    "status": b.status, "status_name": services.BID_STATUSES[b.status],
                    "hold": services.bid_hold(t, b.status)})
    amount, count = services.committed(db, profile.id)
    return {"bids": out, "committed": amount, "committed_count": count, "statuses": services.BID_STATUSES}


@router.get("/tenders/{tender_id}/raw", response_class=PlainTextResponse, summary="Исходный файл извещения")
def tender_raw(tender_id: int, db: Session = Depends(get_db)):
    row = db.get(Tender, tender_id)
    if row is None:
        raise HTTPException(404, "Закупка не найдена")
    return row.raw_content or ""


# ---------- оценка ----------

class ScoreRequest(BaseModel):
    tender_id: int | None = None
    tender: CanonicalTender | None = None
    profile_id: int | None = None
    company_inn: str | None = None
    preferences: Preferences | None = None


@router.post("/score", summary="Оценить закупку: по id или передав каноническую модель, с профилем или настройками")
def score(body: ScoreRequest, db: Session = Depends(get_db)):
    if body.tender is not None:
        tender = body.tender
    elif body.tender_id is not None:
        row = db.get(Tender, body.tender_id)
        if row is None:
            raise HTTPException(404, "Закупка не найдена")
        tender = services.tender_model(row)
    else:
        raise HTTPException(400, "Передайте tender_id или tender")
    if body.preferences is not None:
        prefs = body.preferences
        company_card = get_company(db, body.company_inn) if body.company_inn else None
    else:
        profile = _profile(db, body.profile_id)
        prefs = Preferences.model_validate(services.current_version(db, profile).preferences)
        company_card = services.profile_company(db, profile)
    return services.compute(db, tender, company_card, prefs)


class WhatIfRequest(BaseModel):
    tender_id: int
    profile_id: int | None = None
    nmck_change_pct: float = Field(0, ge=-90, le=300)
    days_left: float | None = Field(None, ge=0, le=120)


@router.post("/score/whatif", summary="«Что если»: оценка при другой НМЦК или сроке подачи — чувствительность вердикта")
def score_whatif(req: WhatIfRequest, db: Session = Depends(get_db)):
    from datetime import timedelta

    row = db.get(Tender, req.tender_id)
    if row is None:
        raise HTTPException(404, "Закупка не найдена")
    profile = _profile(db, req.profile_id)
    pv = services.current_version(db, profile)
    t = services.tender_model(row).model_copy(deep=True)
    k = 1 + req.nmck_change_pct / 100
    if t.nmck is not None:
        t.nmck *= k
        # обеспечения задаются в % от НМЦК — масштабируются вместе с ней
        if t.app_guarantee_amount is not None:
            t.app_guarantee_amount *= k
        if t.contract_guarantee_amount is not None:
            t.contract_guarantee_amount *= k
    if req.days_left is not None:
        t.submission_deadline = datetime.now(timezone.utc) + timedelta(days=req.days_left)
    result = services.compute(db, t, services.profile_company(db, profile), Preferences.model_validate(pv.preferences),
                              committed=services.committed(db, profile.id, row.id))
    return {"nmck": t.nmck, "result": result}


class PreviewRequest(BaseModel):
    profile_id: int | None = None
    preferences: Preferences
    tender_id: int | None = None


@router.post("/score/preview", summary="Пересчитать ленту с черновыми настройками (без сохранения)")
def score_preview(body: PreviewRequest, db: Session = Depends(get_db)):
    return services.preview(db, _profile(db, body.profile_id), body.preferences, body.tender_id)


# ---------- лента и экспорт ----------

@router.get("/feed", summary="Лента закупок, отсортированная по оценке")
def feed(profile_id: int | None = None, verdict: str | None = Query(None, pattern="^(go|consider|skip)$"),
         q: str | None = None, sort: str = Query("score", pattern="^(score|deadline|nmck)$"),
         limit: int = Query(50, le=500), offset: int = 0, db: Session = Depends(get_db)):
    return services.feed(db, _profile(db, profile_id), verdict, q, sort, limit, offset)


@router.get("/feed/export", summary="Выгрузка ленты в XLSX или CSV")
def feed_export(profile_id: int | None = None, format: str = Query("xlsx", pattern="^(xlsx|csv)$"),
                verdict: str | None = None, db: Session = Depends(get_db)):
    data = services.feed(db, _profile(db, profile_id), verdict, None, "score", 10_000, 0)
    header = ["Номер", "Предмет", "Закон", "Способ", "Заказчик", "Регион", "НМЦК, ₽", "Подача до", "Оценка, %",
              "Вердикт", "Главная причина"]
    rows = [[i["purchase_number"], i["subject"], i["law"], i["procedure_name"], i["customer_name"], i["region_name"],
             i["nmck"], (i["submission_deadline"] or "")[:16].replace("T", " "), i["score"],
             VERDICT_NAMES.get(i["verdict"], i["verdict"]), i["main_reason"]] for i in data["items"]]
    stamp = datetime.now().strftime("%Y%m%d_%H%M")
    if format == "csv":
        buf = io.StringIO()
        writer = csv.writer(buf, delimiter=";")
        writer.writerow(header)
        writer.writerows(rows)
        return StreamingResponse(io.BytesIO(buf.getvalue().encode("utf-8-sig")), media_type="text/csv",
                                 headers={"Content-Disposition": f'attachment; filename="tenders_{stamp}.csv"'})
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Лента"
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    fills = {"Участвовать": "D9F8EA", "Рассмотреть": "FFEDB8", "Не участвовать": "FFDCE3", "Проверить вручную": "E6E2FF"}
    for r in rows:
        ws.append(r)
        fill = fills.get(r[9])
        if fill:
            ws.cell(ws.max_row, 10).fill = PatternFill("solid", fgColor=fill)
    for col, width in zip("ABCDEFGHIJK", (24, 60, 8, 24, 40, 28, 14, 18, 10, 18, 70)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return StreamingResponse(out, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="tenders_{stamp}.xlsx"'})


# ---------- пакетная загрузка ----------

@router.post("/batches", summary="Пакетная загрузка: ZIP или несколько XML/JSON, обработка в очереди")
async def create_batch(files: list[UploadFile] = File(...), profile_id: int | None = Form(None), db: Session = Depends(get_db)):
    profile = _profile(db, profile_id)
    unpacked: list[tuple[str, bytes]] = []
    for f in files:
        content = await _read_upload(f)
        try:
            unpacked += unpack(f.filename or "file", content)
        except Exception:
            raise HTTPException(422, f"Не удалось распаковать {f.filename}")
    if not unpacked:
        raise HTTPException(422, "В загрузке нет файлов .xml или .json")
    batch = Batch(profile_id=profile.id, status="queued", total=len(unpacked), errors=[], tender_ids=[])
    db.add(batch)
    db.flush()
    for name, content in unpacked:
        db.add(BatchFile(batch_id=batch.id, filename=name[:255], content=content.decode("utf-8", errors="replace")))
    db.commit()
    mode = dispatch_batch(batch.id)
    audit("batch_created", f"Пакет № {batch.id}: {batch.total} файлов, очередь {mode}", batch_id=batch.id, total=batch.total)
    return {"batch_id": batch.id, "total": batch.total, "queue": mode}


@router.get("/batches/{batch_id}", summary="Прогресс пакета и его результаты")
def get_batch(batch_id: int, db: Session = Depends(get_db)):
    batch = db.get(Batch, batch_id)
    if batch is None:
        raise HTTPException(404, "Пакет не найден")
    profile = _profile(db, batch.profile_id)
    pv = services.current_version(db, profile)
    items = []
    if batch.status == "done":
        for tid in batch.tender_ids or []:
            row = db.get(Tender, tid)
            if row:
                items.append(services._feed_item(row, services.score_row(db, row, pv, services.profile_company(db, profile))))
        items.sort(key=lambda x: (services.VERDICT_ORDER.get(x["verdict"], 9), -x["score"]))
    return {"id": batch.id, "status": batch.status, "total": batch.total, "processed": batch.processed,
            "failed": batch.failed, "errors": batch.errors, "items": items}


@router.post("/quick-score", summary="Быстрая оценка по ТЗ: выгрузка ЕИС + ИНН + текст параметров → оценка и рекомендация")
async def quick_score(
    file: UploadFile = File(..., description="Извещение ЕИС: .xml или .json"),
    inn: str | None = Form(None, description="ИНН вашей компании (10 или 12 цифр)"),
    criteria_text: str | None = Form(None, description="Параметры скоринга обычным текстом"),
    profile_id: int | None = Form(None, description="Если текста нет — взять настройки этого профиля"),
    save: bool = Form(True, description="Добавить закупку в ленту"),
    documents: list[UploadFile] = File(default=[], description="Проект контракта, ТЗ: .docx, .pdf, .txt"),
    db: Session = Depends(get_db),
):
    t0 = time.perf_counter()
    timings: dict[str, float] = {}
    content = await _read_upload(file)

    t = time.perf_counter()
    try:
        tender = parse_bytes(content, file.filename or "")
    except ParseError as exc:
        audit("error", f"Быстрая оценка: файл {file.filename} не разобран — {exc}", level="warning")
        raise HTTPException(422, str(exc))
    timings["parse_ms"] = (time.perf_counter() - t) * 1000

    doc_files = [(d.filename or "doc", await _read_upload(d)) for d in documents if d.filename]
    if doc_files:
        from app.docs.contract import attach

        t = time.perf_counter()
        tender = attach(tender, services.read_documents(doc_files, tender.law, tender, (inn or "").strip() or None))
        timings["documents_ms"] = (time.perf_counter() - t) * 1000

    t = time.perf_counter()
    company_card = None
    if inn and inn.strip():
        try:
            company_card = get_company(db, inn.strip())
        except InvalidInn as exc:
            raise HTTPException(400, str(exc))
        except CompanyNotFound as exc:
            raise HTTPException(404, str(exc))
        except EgrulError as exc:
            raise HTTPException(502, f"{exc}. Можно оценить без ИНН или повторить позже.")
    customer = try_get_company(db, tender.customer_inn)
    timings["egrul_ms"] = (time.perf_counter() - t) * 1000

    t = time.perf_counter()
    parsed = None
    if criteria_text and criteria_text.strip():
        parsed = parse_criteria(criteria_text, Preferences(), timeout=settings.llm_parse_timeout_seconds)
        prefs = parsed.preferences
    else:
        profile = _profile(db, profile_id)
        prefs = Preferences.model_validate(services.current_version(db, profile).preferences)
        if company_card is None:
            company_card = services.profile_company(db, profile)
    timings["criteria_ms"] = (time.perf_counter() - t) * 1000

    from app.scoring.engine import evaluate

    result = evaluate(tender, company_card, customer, prefs).model_dump()
    timings["scoring_ms"] = result["elapsed_ms"]

    # Второе мнение LLM — в пределах оставшегося бюджета 10 секунд (с запасом 1 с).
    if settings.llm_review:
        from app.scoring import review as llm_review

        budget = min(settings.llm_review_timeout_seconds, 9.0 - (time.perf_counter() - t0))
        rv = llm_review.review(tender, result, criteria_text, timeout=budget)
        result = llm_review.apply(result, rv)
        timings["review_ms"] = rv.elapsed_ms
        if rv.status != "unavailable":
            audit("llm_review", f"{tender.purchase_number}: {rv.status_name}", duration_ms=rv.elapsed_ms,
                  purchase_number=tender.purchase_number, status=rv.status, changed=rv.changed,
                  algorithm_verdict=rv.algorithm_verdict, final_verdict=rv.final_verdict)

    tender_id = None
    if save:
        row = services.ingest(db, file.filename or "upload.xml", content)
        if tender.contract:
            row.data = {**row.data, "contract": json.loads(tender.contract.model_dump_json()),
                        "advance_percent": tender.advance_percent}
            db.commit()
        tender_id = row.id
    timings["total_ms"] = (time.perf_counter() - t0) * 1000
    timings = {k: round(v, 1) for k, v in timings.items()}

    audit("quick_score", f"{tender.purchase_number}: {result['score']:.0f}% — {result['verdict_name']}",
          duration_ms=timings["total_ms"], purchase_number=tender.purchase_number, inn=inn,
          verdict=result["verdict"], criteria_engine=parsed.engine if parsed else "profile", timings=timings)
    return {
        "tender": tender,
        "tender_id": tender_id,
        "company": company_card,
        "customer": customer,
        "criteria": parsed.to_dict() if parsed else None,
        "preferences": prefs,
        "result": result,
        "timings": timings,
    }


# ---------- сторона заказчика ----------

@router.post("/customer/check", summary="Заказчику: проверить свою закупку — риски ФАС, условия контракта, интерес поставщиков")
async def customer_check(
    file: UploadFile = File(..., description="Извещение ЕИС: .xml или .json"),
    documents: list[UploadFile] = File(default=[], description="Проект контракта, ТЗ"),
    db: Session = Depends(get_db),
):
    from app import customer as cust
    from app.docs.contract import attach

    started = time.perf_counter()
    content = await _read_upload(file)
    try:
        tender = parse_bytes(content, file.filename or "")
    except ParseError as exc:
        raise HTTPException(422, str(exc))
    docs = [(d.filename or "doc", await _read_upload(d)) for d in documents if d.filename]
    if docs:
        tender = attach(tender, services.read_documents(docs, tender.law, tender))
    findings = cust.check_notice(tender)
    interest = cust.supplier_interest(tender)
    summary = {lvl: sum(1 for f in findings if f.severity == lvl) for lvl in ("high", "warn", "info")}
    elapsed = (time.perf_counter() - started) * 1000
    audit("customer_check", f"{tender.purchase_number}: рисков {summary['high']}, замечаний {summary['warn']}",
          duration_ms=elapsed, purchase_number=tender.purchase_number)
    return {"tender": tender, "findings": findings, "summary": summary, "interest": interest,
            "elapsed_ms": round(elapsed, 1)}


class ParticipantsRequest(BaseModel):
    tender: CanonicalTender
    participants: list[dict] = Field(max_length=30)


@router.post("/customer/participants", summary="Заказчику: проверить участников по ИНН (ЕГРЮЛ, МСП, ОКВЭД, демпинг)")
def customer_participants(req: ParticipantsRequest, db: Session = Depends(get_db)):
    from app import customer as cust

    out = []
    for raw in req.participants:
        p = cust.Participant.model_validate(raw)
        try:
            card = get_company(db, p.inn.strip())
        except InvalidInn as exc:
            raise HTTPException(400, f"{p.inn}: {exc}")
        except EgrulError:
            card = None
        out.append(cust.check_participant(req.tender, card, p))
    return {"participants": out}


@router.get("/logs", summary="Журнал событий: загрузки, оценки, ЕГРЮЛ, разбор критериев, ошибки")
def logs(limit: int = Query(50, le=500), event: str | None = None, db: Session = Depends(get_db)):
    stmt = select(AuditLog).order_by(AuditLog.id.desc()).limit(limit)
    if event:
        stmt = stmt.where(AuditLog.event == event)
    return [
        {"id": r.id, "created_at": r.created_at if r.created_at.tzinfo else r.created_at.replace(tzinfo=timezone.utc), "event": r.event, "event_name": EVENT_NAMES.get(r.event, r.event),
         "level": r.level, "message": r.message, "duration_ms": r.duration_ms, "detail": r.detail}
        for r in db.scalars(stmt)
    ]


@router.get("/demo/sample.xml", response_class=PlainTextResponse, summary="Реальное извещение ЕИС для быстрой проверки")
def demo_sample():
    path = next((Path(__file__).resolve().parents[1] / "samples").glob("real_44fz_ef2020_*.xml"))
    return PlainTextResponse(path.read_text(encoding="utf-8"), media_type="application/xml",
                             headers={"Content-Disposition": f'attachment; filename="{path.name}"'})


@router.get("/demo/batch.zip", summary="Демо-пакет извещений со свежими датами — для проверки пакетной загрузки")
def demo_zip():
    import zipfile

    from app.seed.demo import demo_files

    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, content in demo_files():
            zf.writestr(name, content)
        real = Path(__file__).resolve().parents[1] / "samples"
        for path in real.glob("real_44fz_ef2020_*.xml"):
            zf.writestr(path.name, path.read_bytes())
    buf.seek(0)
    return StreamingResponse(buf, media_type="application/zip",
                             headers={"Content-Disposition": 'attachment; filename="demo_notices.zip"'})


# ---------- обратная связь и точность ----------

@router.post("/scores/{score_id}/review", summary="Второе мнение LLM по оценке алгоритма (процент не меняется)")
def score_review(score_id: int, db: Session = Depends(get_db)):
    from app.scoring import review as llm_review

    score = db.get(Score, score_id)
    if score is None:
        raise HTTPException(404, "Оценка не найдена")
    tender = CanonicalTender.model_validate(db.get(Tender, score.tender_id).data)
    pv = db.get(ProfileVersion, score.profile_version_id)
    base = {**score.result, "verdict": score.result.get("review", {}).get("algorithm_verdict", score.result["verdict"])}
    base["verdict_name"] = VERDICT_NAMES[base["verdict"]]
    rv = llm_review.review(tender, base, pv.criteria_text, timeout=max(settings.llm_review_timeout_seconds, 8.0))
    if rv.status == "unavailable":
        return {"review": rv, "result": score.result}
    result = llm_review.apply(base, rv)
    score.result, score.verdict = result, result["verdict"]
    db.commit()
    audit("llm_review", f"{tender.purchase_number}: {rv.status_name}", duration_ms=rv.elapsed_ms,
          purchase_number=tender.purchase_number, status=rv.status, changed=rv.changed,
          algorithm_verdict=rv.algorithm_verdict, final_verdict=rv.final_verdict)
    return {"review": rv, "result": result}


class FeedbackIn(BaseModel):
    correct: bool
    expected_verdict: str | None = Field(None, pattern="^(go|consider|skip)$")
    comment: str | None = Field(None, max_length=1000)


@router.post("/scores/{score_id}/feedback", summary="Отметка «вердикт верный / неверный»")
def feedback(score_id: int, body: FeedbackIn, db: Session = Depends(get_db)):
    if db.get(Score, score_id) is None:
        raise HTTPException(404, "Оценка не найдена")
    db.add(Feedback(score_id=score_id, correct=body.correct, expected_verdict=body.expected_verdict, comment=body.comment))
    db.commit()
    audit("feedback", f"Оценка {score_id}: вердикт {'верный' if body.correct else 'неверный'}", score_id=score_id)
    return {"ok": True}


@router.get("/learning/suggest", summary="Самообучение: какие веса и пороги лучше совпадут с отметками пользователя")
def learning_suggest(profile_id: int | None = None, db: Session = Depends(get_db)):
    from app.learning import suggest

    return suggest(db, _profile(db, profile_id))


@router.get("/accuracy", summary="Точность: эталонный набор + отметки пользователей")
def accuracy(db: Session = Depends(get_db)):
    total = db.scalar(select(func.count()).select_from(Feedback)) or 0
    correct = db.scalar(select(func.count()).select_from(Feedback).where(Feedback.correct.is_(True))) or 0
    report = json.loads(EVAL_REPORT.read_text(encoding="utf-8")) if EVAL_REPORT.exists() else None
    timings = db.execute(select(func.avg(Score.elapsed_ms), func.max(Score.elapsed_ms))).one()
    return {
        "feedback": {"total": total, "correct": correct, "rate": round(correct / total, 3) if total else None},
        "golden": report,
        "timing_ms": {"avg": round(timings[0] or 0, 2), "max": round(timings[1] or 0, 2)},
    }
