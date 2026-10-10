"""Разметка реальных извещений людьми — эталон «самого верного» на чужих данных.

    python -m app.labeling export --n 120     # корпус ЕИС → eval/labeling/to_label.xlsx
    # два человека независимо заполняют колонки «Разметчик 1» и «Разметчик 2»
    python -m app.labeling import eval/labeling/to_label.xlsx   # → eval/golden_real.yaml + согласие разметчиков
    python -m app.evaluation golden_real      # точность системы на реальных извещениях

Система вердикт не подсказывает: в таблице только данные закупки и профиль поставщика.
В эталон попадают только случаи, где разметчики согласны; расхождения — отдельным списком на обсуждение.
Согласие считается каппой Коэна: она показывает, насколько сама задача однозначна для людей.
"""

import argparse
import json
import random
import sys
from datetime import datetime, timezone
from pathlib import Path

import yaml
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

from app.corpus import CORPUS
from app.eis.parser import parse_bytes
from app.reference import textvec

ROOT = Path(__file__).resolve().parents[1] / "eval"
OUT = ROOT / "labeling"
GOLDEN = ROOT / "golden.yaml"
NAMES = {"go": "Участвовать", "consider": "Рассмотреть", "skip": "Не участвовать"}
BY_NAME = {v.lower(): k for k, v in NAMES.items()}


def _spec(t, now: datetime) -> dict:
    spec = {"subject": t.subject, "okpd2": t.okpd2[:5], "method": t.procedure_type, "law": t.law}
    if t.nmck:
        spec["nmck"] = round(t.nmck, 2)
    if t.delivery_place:
        spec["place"] = t.delivery_place[:200]
    if t.submission_deadline:
        dl = t.submission_deadline if t.submission_deadline.tzinfo else t.submission_deadline.replace(tzinfo=timezone.utc)
        spec["days"] = round((dl - now).total_seconds() / 86400, 2)
    for key in ("smp_only", "national_regime", "contract_guarantee_percent", "customer_inn"):
        if getattr(t, key) is not None:
            spec[key] = getattr(t, key)
    if t.app_guarantee_amount and t.nmck:
        spec["app_guarantee_percent"] = round(100 * t.app_guarantee_amount / t.nmck, 2)
    return spec


def export(n: int, seed: int = 7) -> Path:
    profiles = yaml.safe_load(GOLDEN.read_text(encoding="utf-8"))["profiles"]
    usable = {k: p for k, p in profiles.items() if k != "liquidating"}
    vecs = {k: textvec.embed(p["criteria"]) for k, p in usable.items()}
    now = datetime.now(timezone.utc)
    pool = []
    for f in sorted(CORPUS.glob("*.xml")):
        t = parse_bytes(f.read_bytes(), f.name)
        sims = {k: textvec.cosine(textvec.embed(t.text_for_matching()), v) for k, v in vecs.items()}
        best = max(sims, key=sims.get)
        pool.append((sims[best], best, f.stem, t))
    # Пары «закупка × профиль»: больше похожих на профиль (там сложные решения), часть — случайных.
    pool.sort(key=lambda x: -x[0])
    rng = random.Random(seed)
    near = pool[: int(n * 0.75)]
    rest = pool[int(n * 0.75):]
    far = [(s, rng.choice(list(usable)), stem, t) for s, _, stem, t in rng.sample(rest, min(len(rest), n - len(near)))]
    chosen = near + far
    rng.shuffle(chosen)

    OUT.mkdir(parents=True, exist_ok=True)
    cases = {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Разметка"
    head = ["id", "Профиль поставщика", "Предмет закупки", "НМЦК, ₽", "Место поставки", "Способ", "Дней до конца подачи",
            "Только СМП", "Обеспечение контракта, %", "Ссылка на ЕИС", "Разметчик 1", "Разметчик 2", "Комментарий"]
    ws.append(head)
    for i, (_, profile, stem, t) in enumerate(chosen, 1):
        cid = f"real-{i:03d}"
        spec = _spec(t, now)
        cases[cid] = {"profile": profile, "reestr_number": stem, "tender": spec}
        ws.append([cid, usable[profile]["criteria"], t.subject, t.nmck, t.delivery_place, t.procedure_name or t.procedure_type,
                   spec.get("days"), "да" if t.smp_only else "нет", t.contract_guarantee_percent,
                   f"https://zakupki.gov.ru/epz/order/notice/printForm/view.html?regNumber={stem}", None, None, None])
    dv = DataValidation(type="list", formula1='"Участвовать,Рассмотреть,Не участвовать"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K2:L{len(chosen) + 1}")
    for col, width in zip("ABCDEFGHIJKLM", (10, 60, 60, 14, 40, 22, 10, 8, 10, 30, 16, 16, 30)):
        ws.column_dimensions[col].width = width
    xlsx = OUT / "to_label.xlsx"
    wb.save(xlsx)
    (OUT / "cases.yaml").write_text(yaml.safe_dump({"now": now.isoformat(), "cases": cases}, allow_unicode=True, sort_keys=False),
                                    encoding="utf-8")
    return xlsx


def kappa(a: list[str], b: list[str]) -> float:
    n = len(a)
    if not n:
        return 0.0
    po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((a.count(c) / n) * (b.count(c) / n) for c in NAMES)
    return round((po - pe) / (1 - pe), 3) if pe < 1 else 1.0


def import_labels(xlsx: Path) -> dict:
    meta = yaml.safe_load((OUT / "cases.yaml").read_text(encoding="utf-8"))
    golden = yaml.safe_load(GOLDEN.read_text(encoding="utf-8"))
    ws = load_workbook(xlsx).active
    a_all, b_all, agreed, disagreed = [], [], [], []
    for row in ws.iter_rows(min_row=2, values_only=True):
        cid, l1, l2, comment = row[0], row[10], row[11], row[12]
        if not cid or not l1 or not l2:
            continue
        v1, v2 = BY_NAME.get(str(l1).strip().lower()), BY_NAME.get(str(l2).strip().lower())
        if not v1 or not v2:
            continue
        a_all.append(v1)
        b_all.append(v2)
        case = meta["cases"][cid]
        if v1 == v2:
            agreed.append({"id": cid, "profile": case["profile"], "expected": v1,
                           "why": comment or f"Разметка двух человек, ЕИС № {case['reestr_number']}", "tender": case["tender"]})
        else:
            disagreed.append({"id": cid, "reestr_number": case["reestr_number"], "labeler_1": v1, "labeler_2": v2,
                              "comment": comment})
    out = {"now": meta["now"], "profiles": golden["profiles"], "cases": agreed}
    (ROOT / "golden_real.yaml").write_text(
        "# Эталон на реальных извещениях ЕИС: только случаи, где два разметчика независимо согласны.\n"
        + yaml.safe_dump(out, allow_unicode=True, sort_keys=False), encoding="utf-8")
    summary = {"labeled": len(a_all), "agreed": len(agreed), "agreement": round(len(agreed) / len(a_all), 3) if a_all else None,
               "kappa": kappa(a_all, b_all), "disagreements": disagreed}
    (OUT / "agreement.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    return summary


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("cmd", choices=["export", "import"])
    ap.add_argument("path", nargs="?")
    ap.add_argument("--n", type=int, default=120)
    args = ap.parse_args()
    if args.cmd == "export":
        path = export(args.n)
        print(f"Таблица для разметки: {path} — два человека заполняют колонки «Разметчик 1» и «Разметчик 2» независимо.")
        return 0
    s = import_labels(Path(args.path or OUT / "to_label.xlsx"))
    print(f"Размечено {s['labeled']}, согласны {s['agreed']} ({s['agreement']:.0%}), каппа Коэна {s['kappa']}. "
          f"Эталон: eval/golden_real.yaml; расхождения — eval/labeling/agreement.json")
    return 0


if __name__ == "__main__":
    sys.exit(main())
